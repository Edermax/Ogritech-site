from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "auditoria-2026-09-01"
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "relatorio-auditoria-tecnica-ogritech-2026-09-01.docx"

NAVY = "132A3A"
TEAL = "1E6B55"
BLUE = "2878A8"
ORANGE = "C56B2D"
RED = "A33A32"
GOLD = "B58522"
LIGHT = "F2F6F7"
MID = "DCE7EB"
WHITE = "FFFFFF"
TEXT = "1A2B34"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.75)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.78)
sec.right_margin = Inches(0.78)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08
for name, size, color, before, after in [
    ("Title", 25, NAVY, 0, 8),
    ("Heading 1", 17, NAVY, 15, 7),
    ("Heading 2", 13, BLUE, 11, 5),
    ("Heading 3", 11, TEAL, 8, 4),
]:
    st = styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = name != "Title"
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

if "Callout" not in styles:
    callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
else:
    callout = styles["Callout"]
callout.font.name = "Calibri"
callout.font.size = Pt(11)
callout.font.bold = True
callout.font.color.rgb = RGBColor.from_string(NAVY)
callout.paragraph_format.space_before = Pt(6)
callout.paragraph_format.space_after = Pt(8)
callout.paragraph_format.left_indent = Inches(0.18)
callout.paragraph_format.right_indent = Inches(0.18)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_text(cell, text, bold=False, color=TEXT, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def table(headers, rows, widths=None, header_fill=NAVY, font_size=8.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    hdr = t.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        set_cell_text(hdr.cells[i], h, True, WHITE, 8.5)
        shade(hdr.cells[i], header_fill)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, False, TEXT, font_size)
            if ridx % 2 == 1:
                shade(cells[i], LIGHT)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p

def numbered(text, number):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.add_run(f"{number}.  ").bold = True
    p.add_run(text)
    return p

def callout(text, color=MID):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    shade(c, color)
    set_cell_text(c, text, True, NAVY, 11)
    c.paragraphs[0].paragraph_format.space_before = Pt(7)
    c.paragraphs[0].paragraph_format.space_after = Pt(7)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def page_break():
    doc.add_page_break()

def add_page_number(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.extend([fldChar1, instrText, fldChar2])

header = sec.header.paragraphs[0]
header.text = "OGRITECH  |  AUDITORIA TÉCNICA E DE DADOS"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor.from_string(BLUE)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("Relatório de 01/09/2026  •  Página ")
fr.font.size = Pt(8)
fr.font.color.rgb = RGBColor.from_string("667780")
add_page_number(footer)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(76)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("OGRITECH")
r.bold = True; r.font.size = Pt(15); r.font.color.rgb = RGBColor.from_string(TEAL)
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Relatório de Auditoria Técnica,\nde Dados e de Operação")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Site, front-end, Supabase, segurança, qualidade e governança")
r.font.size = Pt(13); r.font.color.rgb = RGBColor.from_string(BLUE)
doc.add_paragraph()
table(["Escopo", "Data-base", "Objetivo"], [["Código + site publicado + Supabase", "1º de setembro de 2026", "Reduzir retrabalho, correções tardias e desperdício operacional"]], [2.2, 1.3, 3.0], TEAL, 9)
callout("Conclusão: a base é viável para piloto controlado. O melhor retorno agora vem de consolidar arquitetura, testes end-to-end, políticas do banco e operação — não de uma reescrita geral.", "E8F4EF")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Preparado como avaliação técnica independente baseada em evidências do repositório e do ambiente publicado.")
r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string("667780")

page_break()
doc.add_heading("1. Resumo executivo", level=1)
doc.add_paragraph("A Ogritech possui uma fundação funcional e acima da média para um produto ainda em pré-operação: autenticação real, isolamento multiempresa com RLS, 36 migrations aplicadas nos dois ambientes, Edge Function autenticada, documentação de continuidade, CI com banco limpo e rotas públicas verificadas sem erros de console. Os 16 testes automatizados locais passaram integralmente.")
doc.add_paragraph("O risco dominante não é falha estrutural imediata; é o custo crescente de alteração. A interface principal está concentrada em arquivos grandes e páginas monolíticas, existem dois pontos de entrada de login, há 73 construções de HTML dinâmico e o CI ainda não cobre jornadas reais no navegador. No banco de produção, o Performance Advisor retorna 105 observações: 8 políticas com auth reavaliado por linha, 17 conjuntos de políticas permissivas sobrepostas, 7 chaves estrangeiras sem índice e 73 índices ainda não utilizados.")
callout("Decisão recomendada: manter a tecnologia atual, congelar novas frentes por um ciclo curto de estabilização e executar um plano de 90 dias orientado por risco e aprendizado.", "FFF3D9")

doc.add_heading("Avaliação geral", level=2)
table(["Dimensão", "Situação", "Nota", "Leitura"], [
    ["Funcionalidade pública", "Boa", "8/10", "Rotas principais respondem, estados de erro são compreensíveis e não houve erro de console."],
    ["Segurança de aplicação", "Boa com lacunas", "7/10", "RLS, CSP em meta e Edge Function protegida; falta endurecimento de headers e proteção contra senhas vazadas."],
    ["Banco e isolamento", "Boa", "8/10", "RLS nas tabelas públicas e migrations consistentes; há dívida de performance e defesa em profundidade no schema private."],
    ["Manutenibilidade", "Atenção", "5/10", "Arquivos centrais grandes, HTML dinâmico frequente, duplicação de entrada e ausência de tipagem gerada."],
    ["Qualidade e testes", "Boa base", "7/10", "CI e pgTAP fortes; faltam E2E, acessibilidade, regressão visual e orçamento de desempenho."],
    ["Dados e gestão", "Inicial", "5/10", "Há dados operacionais e auditoria, mas ainda não existe camada canônica de métricas do negócio."],
    ["Operação e jurídico", "Bloqueador parcial", "5/10", "Backup, restauração e monitoramento existem; contratos, SLAs e revisão jurídica permanecem abertos."],
], [1.45, 1.25, 0.65, 3.15], NAVY, 8.2)

doc.add_heading("Cinco ações que mais evitam retrabalho", level=2)
for n, text in enumerate([
    "Criar uma camada de módulos de domínio no front-end e reduzir gradualmente script.js, style.css, admin.js e cliente.js; evitar uma reescrita de uma vez.",
    "Adicionar testes end-to-end das seis jornadas críticas e tornar esses testes bloqueadores de publicação.",
    "Consolidar as políticas RLS duplicadas e corrigir os oito casos de auth initplan antes do crescimento de volume.",
    "Definir um único contrato de rotas, um único login e uma política de versões/depreciação para páginas públicas.",
    "Fechar os itens jurídicos, fiscais, SLAs e responsabilidades operacionais antes de ampliar clientes ou dados reais.",
], 1): numbered(text, n)

page_break()
doc.add_heading("2. Escopo, método e evidências", level=1)
doc.add_paragraph("A análise foi executada sem alterar produção. Foram usados quatro tipos de evidência: código versionado, testes automatizados, inspeção read-only dos projetos Supabase e navegação das páginas públicas publicadas.")
table(["Fonte", "Cobertura", "Evidência observada"], [
    ["Repositório local", "HTML, CSS, JavaScript, SQL, CI e documentação", "140 arquivos rastreados; 19 HTML, 23 JS, 43 SQL e 17 documentos Markdown no conjunto analisado."],
    ["Testes locais", "Estático, segurança, monitoramento e banco", "npm run validate: 16 verificações estáticas, 10 páginas, 36 migrations e 16 testes aprovados."],
    ["Site publicado", "Home, solução, contato, login, demonstração, páginas públicas, legais e 404", "Rotas principais responderam e não apresentaram erros de console na verificação."],
    ["Supabase produção", "Advisors, tabelas, migrations e Edge Function", "Projeto saudável em sa-east-1, PostgreSQL 17.6, 36 migrations e platform-users ativa com verify_jwt=true."],
    ["Supabase staging", "Paridade estrutural", "Projeto saudável, mesmas 36 migrations e mesma Edge Function protegida."],
], [1.35, 2.0, 3.2], NAVY, 8.3)
doc.add_paragraph("Limitações: não foram executadas operações de escrita, testes com credenciais reais, carga de alto volume, pentest externo, análise jurídica nem validação de entrega de e-mail. Os dados de produção ainda são esparsos e não sustentam conclusões sobre retenção, receita, conversão ou comportamento de clientes.")

doc.add_heading("3. Arquitetura atual", level=1)
doc.add_paragraph("A solução é uma aplicação web multipágina baseada em HTML/CSS/JavaScript estáticos, com Supabase Auth, Postgres/RLS, RPCs e uma Edge Function administrativa. A entrega pública passa por Cloudflare sobre origem do GitHub Pages. O desenho é adequado para o estágio atual e reduz custo operacional, mas exige disciplina para evitar que a lógica de interface continue concentrada em arquivos globais.")
table(["Camada", "Componentes", "Risco principal"], [
    ["Aquisição", "Home, Agenda Online, Contato e Demonstrações", "Múltiplas páginas e estilos podem divergir sem um design system versionado."],
    ["Acesso", "Login, recuperação de senha, painel, cliente e admin", "Duplicação de entrada e concentração de permissões/jornadas em arquivos grandes."],
    ["Páginas do cliente", "Agendamento, página comercial, cardápio, pedido e proposta", "Contratos de URL com slug/token precisam de testes de compatibilidade e expiração."],
    ["Dados", "Postgres 17, RLS, RPCs, auditoria e billing", "Políticas sobrepostas e índices exigem acompanhamento antes de escala."],
    ["Operação", "CI, staging, monitor, backup e restauração", "Aprovação do go-live e responsabilidades de negócio ainda não estão formalizadas."],
], [1.2, 2.55, 2.8], TEAL, 8.2)

page_break()
doc.add_heading("4. Achados técnicos prioritários", level=1)
doc.add_heading("P0 — fechar antes de ampliar o piloto", level=2)
table(["Achado", "Impacto", "Correção recomendada", "Aceite"], [
    ["Pendências jurídicas e operacionais", "Risco contratual, fiscal, LGPD e de atendimento; pode transformar incidente técnico em problema de negócio.", "Concluir razão social/CNPJ, revisão jurídica, canais, SLA, escala de incidentes, emissão fiscal e gateway quando aplicável.", "Checklist de lançamento assinado por responsáveis técnico e de negócio."],
    ["Go-live sem registro final", "Dificulta auditoria, rollback e responsabilização.", "Preencher projeto, commit, data/hora, responsáveis e resultado do smoke test em cada release.", "Registro imutável anexado ao release."],
], [1.4, 1.65, 2.35, 1.1], RED, 7.7)

doc.add_heading("P1 — executar no próximo ciclo", level=2)
table(["Achado", "Evidência", "Risco", "Ação"], [
    ["Sem proteção contra senhas vazadas", "Security Advisor de produção e staging: 1 warning.", "Credenciais comprometidas podem ser reutilizadas.", "Ativar Leaked Password Protection e testar cadastro/troca de senha."],
    ["Headers incompletos", "Produção entrega HSTS e nosniff, mas não observamos CSP em header nem proteção frame-ancestors/X-Frame-Options.", "Meta CSP não cobre todos os controles e não protege adequadamente contra framing.", "Publicar CSP como header no proxy/CDN, definir frame-ancestors 'none' e revisar Referrer-Policy e Permissions-Policy."],
    ["CI sem jornadas de navegador", "CI executa npm validate e pgTAP, mas não executa browser E2E.", "Alterações podem quebrar login, redirects, formulários e responsividade sem falhar o pipeline.", "Adicionar Playwright em staging com smoke tests das jornadas críticas."],
    ["Front-end monolítico", "script.js: 88 KB/1.539 linhas; style.css: 59 KB/1.506 linhas; 73 usos de HTML dinâmico.", "Aumenta acoplamento, regressões e tempo de revisão.", "Extrair módulos por domínio, componentes de UI e utilitários seguros; migrar por fatias."],
    ["Performance RLS", "8 auth_rls_initplan e 17 conjuntos de políticas permissivas sobrepostas.", "Cada linha pode reavaliar funções; políticas extras elevam custo e complexidade.", "Consolidar políticas preservando a matriz de autorização e trocar auth.fn() por (select auth.fn()) quando seguro."],
    ["Chaves estrangeiras sem índice", "7 alertas do Performance Advisor.", "Deleções, joins e filtros podem degradar com volume.", "Adicionar índices após confirmar padrões de consulta e medir planos."],
], [1.35, 1.75, 1.55, 2.0], ORANGE, 7.4)

doc.add_heading("P2 — reduzir custo futuro", level=2)
for text in [
    "Unificar /login/ e login.html; manter um redirect temporário e registrar data de remoção.",
    "Formatar e modularizar cardapio/index.html e pagina/index.html, hoje condensados em uma linha, para melhorar revisão e diffs.",
    "Gerar tipos TypeScript do Supabase e introduzir JSDoc/TypeScript nas bordas de domínio antes de converter toda a aplicação.",
    "Criar tokens de design e componentes básicos compartilhados para evitar divergência entre home, páginas públicas e painel.",
    "Revisar os 73 índices não utilizados somente após janela suficiente de métricas; não removê-los automaticamente em pré-operação.",
]: bullet(text)

page_break()
doc.add_heading("5. Banco de dados, Supabase e segurança", level=1)
doc.add_paragraph("Os dois projetos estão ACTIVE_HEALTHY e executam PostgreSQL 17.6.1.155. Produção e staging possuem as mesmas 36 migrations, reduzindo risco de drift. A Edge Function platform-users está ativa e exige JWT nos dois ambientes. Todas as tabelas do schema public listadas pelo conector estão com RLS habilitada.")

doc.add_heading("Leitura correta dos alertas do schema private", level=2)
doc.add_paragraph("O conector sinalizou private.public_booking_attempts e private.platform_legal_identity sem RLS. Como o schema private não deve estar exposto pela Data API, isso não equivale automaticamente a exposição pública. Ainda assim, RLS é recomendada como defesa em profundidade e os privilégios de schema/tabela devem ser auditados explicitamente. Não se deve simplesmente ativar RLS em produção sem validar as funções que dependem dessas tabelas.")
callout("SQL proposto para avaliação em staging — não executado: ALTER TABLE private.public_booking_attempts ENABLE ROW LEVEL SECURITY; ALTER TABLE private.platform_legal_identity ENABLE ROW LEVEL SECURITY. Antes, confirmar grants, chamadas SECURITY DEFINER e testes funcionais.", "FFF3D9")

doc.add_heading("Alertas atuais de produção", level=2)
table(["Tipo", "Quantidade", "Prioridade", "Tratamento"], [
    ["Leaked password protection", "1", "Alta", "Ativar no Auth e validar fluxos de senha."],
    ["auth_rls_initplan", "8", "Alta antes de escala", "Otimizar políticas e repetir pgTAP/advisors."],
    ["Políticas permissivas múltiplas", "17", "Alta antes de escala", "Consolidar por ação/papel sem ampliar acesso."],
    ["FKs sem índice", "7", "Média", "Priorizar joins e tabelas de maior crescimento."],
    ["Índices não utilizados", "73", "Baixa agora", "Observar por 60–90 dias antes de remover."],
], [2.0, 0.8, 1.4, 2.3], NAVY, 8.3)

doc.add_heading("Mudança de plataforma a acompanhar", level=2)
doc.add_paragraph("O changelog do Supabase informa que novas tabelas deixarão de ser expostas automaticamente pela Data API, com aplicação ampla prevista para 30/10/2026. O repositório já declara grants explícitos e RLS, o que é positivo. A ação preventiva é incluir em cada migration um teste que valide: schema exposto, grants para anon/authenticated, RLS ativa e políticas mínimas para cada nova tabela pública.")

doc.add_heading("6. Qualidade do front-end e experiência", level=1)
doc.add_paragraph("A inspeção das páginas públicas confirmou títulos, H1, robots e estados de erro coerentes. Home, Agenda Online, Contato, Login, Demonstrações, Agendamento, Página Comercial, Cardápio, Privacidade, Termos e 404 carregaram sem erros de console. As páginas dependentes de empresa exibem mensagens compreensíveis quando o slug está ausente.")
table(["Ponto forte", "Por que importa"], [
    ["Robots noindex nas áreas privadas e páginas parametrizadas", "Reduz indexação de conteúdo operacional e URLs incompletas."],
    ["Estados vazios nas páginas públicas", "Evita telas quebradas quando empresa, cardápio ou página não estão configurados."],
    ["CSP e versões fixadas do supabase-js", "Reduz superfície de scripts, apesar da necessidade de levar CSP ao header."],
    ["404 orientada para recuperação", "Mantém o usuário em uma jornada útil."],
], [2.45, 4.05], TEAL, 8.5)

doc.add_heading("Correções de experiência que evitam retrabalho", level=2)
for text in [
    "Definir estados padrão para loading, vazio, erro recuperável, indisponível e sucesso em todos os módulos.",
    "Criar contrato único de mensagens e validação; hoje cada página pode evoluir de forma independente.",
    "Adicionar testes em 360 px, 768 px e 1440 px, teclado, foco visível e leitor de tela para formulários críticos.",
    "Substituir innerHTML por DOM seguro ou templates com escape centralizado; manter testes com payloads maliciosos.",
    "Adicionar orçamento de performance: LCP, CLS, INP, peso JS/CSS e limite de requisições por página.",
]: bullet(text)

doc.add_heading("7. Estratégia de testes e qualidade", level=1)
doc.add_paragraph("A combinação atual de validação estática, testes Node e pgTAP é uma base sólida. O próximo ganho não vem de duplicar testes unitários, mas de cobrir as integrações que hoje dependem de ambiente, browser, redirects, autenticação e dados.")
table(["Camada", "Estado atual", "Próximo gate"], [
    ["Estática", "16 scripts, 10 páginas e 36 migrations verificados", "Manter e incluir contrato de rotas e HTML válido."],
    ["Node", "16 testes aprovados", "Adicionar testes de utilitários de formatação, escape e transições de estado."],
    ["Banco", "pgTAP em banco limpo no CI", "Adicionar testes de regressão dos advisors e grants de novas tabelas."],
    ["E2E", "Não identificado no CI", "Playwright contra staging em cada PR de risco."],
    ["Visual/a11y", "Não identificado no CI", "Snapshots seletivos, axe e navegação por teclado."],
    ["Performance", "Monitor operacional existente", "Lighthouse/Web Vitals com orçamento e tendência."],
], [1.2, 2.65, 2.65], NAVY, 8.2)

doc.add_heading("Seis jornadas E2E obrigatórias", level=2)
for n, text in enumerate([
    "Login e redirecionamento por papel: owner/admin/employee/client/platform admin.",
    "Recuperação de senha e retorno ao login.",
    "Criação, alteração e cancelamento de agendamento, incluindo conflito simultâneo.",
    "Agendamento público completo e gestão via token.",
    "Lead → orçamento → proposta, com autorização e expiração de token.",
    "Cardápio → pedido → acompanhamento, incluindo falha de rede e reenvio idempotente.",
], 1): numbered(text, n)

doc.add_heading("8. Engenharia de dados e gestão do negócio", level=1)
doc.add_paragraph("O banco já registra entidades operacionais, faturamento da plataforma, auditoria, privacidade e transições de agendamento. Porém, o produto ainda não possui uma camada analítica canônica. Sem ela, cada dashboard futuro tende a calcular métricas de forma diferente e gerar retrabalho em decisões, relatórios e cobrança.")

doc.add_heading("Camada mínima de métricas", level=2)
table(["Métrica", "Definição inicial", "Fonte", "Guardrail"], [
    ["Empresas ativas", "Negócios com assinatura ativa e ao menos uma atividade válida no período", "saas_clients + subscriptions + eventos", "Excluir demonstrações e contas internas."],
    ["Ativação", "Empresa que configurou serviços, equipe e realizou o primeiro agendamento real", "settings/services/employees/appointments", "Janela explícita de 14 ou 30 dias."],
    ["Uso semanal", "Empresas com operação significativa na semana", "eventos de domínio", "Não contar apenas login."],
    ["Conversão de agenda", "Agendamentos concluídos ÷ tentativas válidas", "public_booking_attempts + appointments", "Separar indisponibilidade, abandono e bloqueio."],
    ["No-show", "Agendamentos marcados como falta ÷ agendamentos esperados", "status events", "Excluir cancelamentos prévios."],
    ["Receita da plataforma", "Pagamentos líquidos de estornos por competência", "invoices/payments/refunds", "Definir competência e caixa separadamente."],
    ["SLA de suporte", "Tempo de primeira resposta e resolução", "futura ferramenta de tickets", "Separar severidade e horário útil."],
], [1.25, 2.5, 1.65, 1.35], TEAL, 7.4)

doc.add_heading("Recomendações de modelagem analítica", level=2)
for text in [
    "Criar eventos de domínio imutáveis para ações relevantes; não depender de logs de console ou da tela.",
    "Separar dados de demonstração com flag/tenant explícito e excluí-los por padrão dos relatórios.",
    "Documentar grão, janela, denominador, exclusões e proprietário de cada KPI antes de criar dashboards.",
    "Criar views analíticas security_invoker ou em schema não exposto, com testes de isolamento multiempresa.",
    "Definir retenção, anonimização e acesso para eventos de produto conforme LGPD e necessidade operacional.",
]: bullet(text)

doc.add_heading("9. Plano recomendado de 90 dias", level=1)
table(["Prazo", "Entregas", "Responsável sugerido", "Critério de saída"], [
    ["0–7 dias", "Fechar checklist de go-live; ativar proteção de senha; definir headers; registrar release; congelar novas features de risco.", "Fundador + técnico + jurídico", "Piloto com responsabilidades, rollback e comunicação definidos."],
    ["8–30 dias", "Playwright das jornadas críticas; consolidar login; corrigir 8 initplans; planejar políticas sobrepostas; índices das FKs prioritárias.", "Full stack + dados", "CI bloqueia regressões e advisors críticos/altos têm plano aprovado."],
    ["31–60 dias", "Extrair módulos de agenda, clientes, serviços e billing; centralizar templates/escape; gerar tipos Supabase; design tokens.", "Full stack", "Mudanças por domínio não exigem editar arquivos globais extensos."],
    ["61–90 dias", "Camada canônica de eventos/KPIs; performance budget; acessibilidade; revisão de índices com dados reais.", "Dados + produto + full stack", "Primeiro scorecard operacional reproduzível e auditável."],
], [0.85, 3.1, 1.25, 1.8], NAVY, 7.6)

doc.add_heading("Ordem que evita desperdício", level=2)
for n, text in enumerate([
    "Não começar por framework novo. Primeiro crie testes de caracterização e contratos de jornada.",
    "Não remover índices por alerta de uso antes de existir tráfego representativo.",
    "Não consolidar RLS sem matriz de papéis e testes negativos de acesso cruzado.",
    "Não criar dashboard antes de definir métricas e separar dados demonstrativos.",
    "Não abrir novos módulos comerciais antes de estabilizar padrões de UI, permissões e auditoria.",
], 1): numbered(text, n)

doc.add_heading("10. Padrões de engenharia para novas mudanças", level=1)
table(["Decisão", "Regra proposta"], [
    ["Arquitetura", "Toda feature deve ter domínio, entradas, saídas, permissões e dono definidos; usar ADR para decisões irreversíveis."],
    ["Banco", "Uma migration por mudança coerente, aplicada em staging, testada com pgTAP e advisors antes de produção."],
    ["Segurança", "Negar por padrão; autorização no banco; tokens mínimos; nenhuma regra de acesso baseada apenas no front-end."],
    ["Frontend", "Sem HTML dinâmico sem escape; estados padrão; componentes compartilhados; compatibilidade de rota documentada."],
    ["Dados", "Eventos com schema versionado; KPIs com definição, proprietário e teste; demonstração excluída por padrão."],
    ["Release", "Commit imutável, smoke test, backup/rollback, observabilidade e responsável de plantão."],
    ["Qualidade", "Teste mais próximo do risco: unitário para regra, integração para contrato e E2E para jornada."],
], [1.25, 5.25], TEAL, 8.4)

doc.add_heading("11. Conclusão", level=1)
doc.add_paragraph("A Ogritech não precisa recomeçar. O produto já reúne os elementos difíceis — autenticação, isolamento, operação de banco, auditoria, staging, backup e testes estruturais. O risco agora é crescer em superfície antes de reduzir o acoplamento e formalizar a operação.")
doc.add_paragraph("A recomendação é um ciclo curto de estabilização com gates claros, seguido de modularização incremental. Isso preserva o investimento atual, reduz o custo de cada nova mudança e cria condições para que produto, dados e operação evoluam sem múltiplas correções tardias.")
callout("Próxima decisão: aprovar o plano 0–30 dias e nomear um responsável por cada entrega. Só depois retomar novas funcionalidades comerciais.", "E8F4EF")

doc.add_heading("Apêndice A — Inventário de evidências", level=1)
table(["Evidência", "Local / referência"], [
    ["Código e arquitetura", "README.md; index.html; script.js; style.css; admin.js; cliente.js; painel/index.html; supabase-config.js"],
    ["CI e testes", ".github/workflows/validate.yml; tests/*.mjs; supabase/tests/database/*.sql"],
    ["Operação", "docs/CHECKLIST_LANCAMENTO.md; docs/OPERACAO_PRODUCAO.md; docs/PLANO_CONTINUIDADE_PILOTO.md"],
    ["Banco", "supabase/migrations/*.sql; Supabase Advisors e inventário read-only de 01/09/2026"],
    ["Site publicado", "https://ogritech.com.br e rotas públicas verificadas em 01/09/2026"],
    ["Documentação externa", "Supabase Changelog; Password Security; RLS; Database Linter"],
], [1.65, 4.85], NAVY, 8.4)

doc.add_heading("Apêndice B — Referências externas", level=1)
for text in [
    "Supabase Changelog: https://supabase.com/changelog",
    "Password security: https://supabase.com/docs/guides/auth/password-security",
    "Row Level Security: https://supabase.com/docs/guides/database/postgres/row-level-security",
    "Database linter: https://supabase.com/docs/guides/database/database-linter",
    "Securing the Data API: https://supabase.com/docs/guides/api/securing-your-api",
]: bullet(text)

doc.core_properties.title = "Relatório de Auditoria Técnica, de Dados e de Operação — Ogritech"
doc.core_properties.subject = "Avaliação do site, código, Supabase, dados, segurança e operação"
doc.core_properties.author = "OpenAI Codex"
doc.core_properties.keywords = "Ogritech, auditoria, full stack, engenharia de dados, Supabase, segurança"
doc.save(DOCX)
print(DOCX)
